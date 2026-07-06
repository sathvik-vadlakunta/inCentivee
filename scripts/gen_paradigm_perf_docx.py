#!/usr/bin/env python3
"""Generate the Paradigm Experts Webflow performance tickets as a polished .docx
for handoff to the client's Webflow developer.

    python3 scripts/gen_paradigm_perf_docx.py [output.docx]
"""

import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x1F, 0x29, 0x37)
MAROON = RGBColor(0x7A, 0x1F, 0x2B)   # Paradigm brand maroon
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
MUTED = RGBColor(0x64, 0x74, 0x8B)

PRIO_COLOR = {"P0": RED, "P1": RGBColor(0xD9, 0x77, 0x06), "P2": MUTED}


def _set_cell(cell, text, *, bold=False, color=None, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def add_metric_table(doc, rows, headers):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        _set_cell(table.rows[0].cells[i], h, bold=True, size=9.5)
    for row in rows:
        cells = table.add_row().cells
        for i, (val, color) in enumerate(row):
            _set_cell(cells[i], val, color=color, bold=(i == 0))
    return table


def ticket(doc, tid, title, prio, owner, problem, tasks, acceptance):
    h = doc.add_heading(level=2)
    run = h.add_run(f"{tid} — {title}")
    run.font.color.rgb = NAVY
    meta = doc.add_paragraph()
    pr = meta.add_run(f"{prio}")
    pr.bold = True
    pr.font.color.rgb = PRIO_COLOR.get(prio, MUTED)
    pr.font.size = Pt(10)
    o = meta.add_run(f"   ·   Owner: {owner}")
    o.font.size = Pt(10)
    o.font.color.rgb = MUTED

    pp = doc.add_paragraph()
    pp.add_run("Problem.  ").bold = True
    pp.add_run(problem)

    doc.add_paragraph("Tasks", style="Heading 4")
    for t in tasks:
        doc.add_paragraph(t, style="List Bullet")

    ap = doc.add_paragraph()
    r = ap.add_run("✓ Acceptance:  ")
    r.bold = True
    r.font.color.rgb = GREEN
    ap.add_run(acceptance)
    doc.add_paragraph()


def build(path):
    doc = Document()
    # base font
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    title = doc.add_heading(level=0)
    tr = title.add_run("Paradigm Experts — Site Performance Tickets")
    tr.font.color.rgb = MAROON
    sub = doc.add_paragraph()
    sr = sub.add_run("paradigmexperts.com (Webflow)  ·  Prepared by PracticeRank  ·  2026-06-29")
    sr.font.color.rgb = MUTED
    sr.font.size = Pt(10)
    sr.italic = True

    doc.add_heading("Why this matters", level=1)
    doc.add_paragraph(
        "Mobile Lighthouse Performance = 51/100. SEO, Accessibility, and Best-Practices are already "
        "96–100, so page speed is the one thing holding the site back — and a 6.8-second load actively "
        "suppresses rankings on the “near me” buyer-intent searches we're working to win (sell gold, "
        "coin & jewelry buyers, diamond buyers, etc.)."
    )

    doc.add_heading("Baseline (mobile, throttled) — re-measure against the targets", level=1)
    add_metric_table(
        doc,
        [
            [("Largest Contentful Paint (LCP)", None), ("6.8 s", RED), ("< 2.5 s", GREEN)],
            [("First Contentful Paint (FCP)", None), ("6.0 s", RED), ("< 1.8 s", GREEN)],
            [("Time to Interactive (TTI)", None), ("9.2 s", RED), ("< 3.8 s", GREEN)],
            [("Total Blocking Time (TBT)", None), ("440 ms", PRIO_COLOR["P1"]), ("< 200 ms", GREEN)],
            [("Cumulative Layout Shift (CLS)", None), ("0  ✓", GREEN), ("< 0.1", GREEN)],
        ],
        ["Metric", "Now", "Target (Good)"],
    )
    p = doc.add_paragraph()
    p.add_run("Lighthouse-flagged opportunities:  ").bold = True
    p.add_run("multiple redirects ≈ 780 ms  ·  unused JavaScript ≈ 450 ms / 284 KB  ·  unused CSS ≈ 14 KB.")
    v = doc.add_paragraph()
    v.add_run("How to verify any ticket:  ").bold = True
    v.add_run(
        "run pagespeed.web.dev on https://paradigmexperts.com/ (Mobile tab), test on a Webflow staging "
        "subdomain first if possible, and compare the metric named in the ticket. CLS is already perfect — "
        "do not let any change introduce layout shift (always set image dimensions)."
    )

    doc.add_heading("Tickets", level=1)

    ticket(
        doc, "TICKET-1", "Optimize the hero / LCP image", "P0", "Webflow dev",
        "LCP is 6.8 s. The largest above-the-fold image is loading too late and too heavy. "
        "This is the single biggest win.",
        [
            "Identify the LCP element (PageSpeed → “Largest Contentful Paint element”) — almost certainly the homepage hero.",
            "If the hero is a background image on a div, switch it to a real Webflow Image element so Webflow generates responsive srcset + serves WebP via its CDN. (Background images don't get responsive variants.)",
            "Re-export the source at its real display size (don't ship a 3000px image into a 1200px slot) and compress before upload.",
            "Set the hero image Loading = “eager” (Webflow Image settings); leave every other image “lazy”.",
            "Preload the hero in Page Settings → Custom Code (Head): <link rel=\"preload\" as=\"image\" href=\"HERO_WEBP_URL\" fetchpriority=\"high\">",
            "Add Custom Attribute on the hero image: name fetchpriority, value high.",
        ],
        "Mobile LCP < 2.5 s; hero served as WebP at ~display resolution; no CLS regression.",
    )
    ticket(
        doc, "TICKET-2", "Lazy-load below-the-fold images + set explicit dimensions", "P1", "Webflow dev",
        "Off-screen images compete for bandwidth during first paint and inflate FCP/LCP.",
        [
            "Every below-the-fold image → Loading = “lazy” (verify; convert any below-fold background images to Image elements).",
            "Ensure every image has explicit width/height (or aspect-ratio) so nothing reflows — protects the perfect CLS = 0.",
        ],
        "Only the hero loads eagerly; “Defer offscreen images” passes; CLS stays < 0.1.",
    )
    ticket(
        doc, "TICKET-3", "Reduce / defer JavaScript", "P1", "Webflow dev",
        "~284 KB of unused JS and ~450 ms of wasted execution; TBT 440 ms and TTI 9.2 s are JS-bound.",
        [
            "Audit Site Settings → Custom Code and every page/element Embed for third-party scripts (chat, pixels, analytics, old tag managers). Remove unused; defer the rest (defer attribute or load on interaction).",
            "Remove unused Webflow Interactions (IX2) — each unused interaction still ships JS.",
            "Consolidate duplicate analytics (only one GA/GTM instance).",
            "Move non-critical custom <script> from <head> to end of <body> / add defer.",
            "Note: Webflow always ships jQuery + webflow.js — that floor can't be removed without leaving Webflow. Eliminate everything on top of that baseline.",
        ],
        "“Reduce unused JavaScript” savings < 100 ms; mobile TBT < 200 ms.",
    )
    ticket(
        doc, "TICKET-4", "Fonts: preconnect + font-display: swap", "P2", "Webflow dev",
        "Web fonts block first paint, hurting FCP (6.0 s).",
        [
            "If using Google/Adobe fonts, add to Head: <link rel=\"preconnect\" href=\"https://fonts.gstatic.com\" crossorigin> (and the Adobe host if used).",
            "Ensure font-display: swap (Webflow applies this for uploaded fonts; for custom @font-face in embeds, add font-display: swap;).",
            "Drop unused font weights/styles from the project.",
        ],
        "No “Ensure text remains visible during webfont load” flag; FCP improves toward < 1.8 s.",
    )
    ticket(
        doc, "TICKET-5", "Collapse the redirect chain", "P1", "DNS/Cloudflare (PracticeRank) + dev to confirm",
        "Lighthouse flags ~780 ms lost to “multiple redirects” — likely http→https→www or a trailing-slash hop.",
        [
            "Pick ONE canonical host (recommend https://paradigmexperts.com, no www) and be consistent.",
            "Webflow → Site Settings → Publishing → set the chosen host as the Default domain.",
            "Ensure DNS/Cloudflare does the canonical redirect in a single 301 (not http→https→www = two hops). One Cloudflare Redirect Rule.",
            "Re-test: curl -sIL https://paradigmexperts.com should show at most one 301 before 200.",
        ],
        "≤ 1 redirect to reach the final URL; “Avoid multiple page redirects” passes.",
    )
    ticket(
        doc, "TICKET-6", "Cloudflare edge optimization", "P1", "DNS/Cloudflare (PracticeRank)",
        "Slow TTFB/FCP. Cloudflare in front of Webflow cuts first-byte and offloads images. "
        "(Skip / adapt if the domain isn't proxied through Cloudflare yet — coordinate with the team.)",
        [
            "Confirm the domain is proxied through Cloudflare (orange cloud) in front of Webflow hosting.",
            "Enable Brotli, Early Hints, and Tiered Cache.",
            "Enable Polish (WebP/AVIF + lossy) for images (Pro plan).",
            "Add a Cache Rule to edge-cache static assets; be careful caching HTML — respect Webflow's cache headers so re-publishes invalidate. Test before/after.",
            "Heads-up: Cloudflare Auto Minify was deprecated — do minification at the source instead (see TICKET-7).",
        ],
        "TTFB drops measurably; images served as AVIF/WebP from the edge; no stale content after a Webflow re-publish.",
    )
    ticket(
        doc, "TICKET-7", "Confirm Webflow publish-time minification is ON", "P2", "Webflow dev",
        "Webflow can minify HTML/CSS/JS at publish — make sure it's enabled.",
        [
            "Webflow → Site Settings → Publishing → enable Minify HTML / CSS / JS. Re-publish.",
        ],
        "Published HTML/CSS/JS are minified.",
    )

    doc.add_heading("Expected outcome", level=1)
    doc.add_paragraph(
        "TICKET-1 alone should move Performance from 51 into the 70s. With 2–7 done, expect mobile ≈ 80–88 "
        "and desktop ≈ 90+. A guaranteed 90+ on mobile is hard on Webflow (the jQuery + webflow.js baseline "
        "is the ceiling); if that's a hard requirement, the durable fix is porting to our Astro + Cloudflare "
        "Pages platform (separate proposal). Re-run PageSpeed after each P0/P1 ticket and log the new score."
    )

    doc.add_heading("Sign-off log", level=1)
    t = add_metric_table(
        doc,
        [[("2026-06-29", None), ("baseline", None), ("51", RED), ("6.8 s", RED), ("starting point", MUTED)]],
        ["Date", "Ticket", "Mobile Perf", "LCP", "Notes"],
    )

    doc.save(path)
    print(f"wrote {path}")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "docs/Paradigm-Experts-Performance-Tickets.docx"
    build(out)
