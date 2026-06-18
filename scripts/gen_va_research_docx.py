#!/usr/bin/env python3
"""Generate the VA-handoff research plan for Lever A (digital PR / earned mentions).

Outputs a .docx the operator can hand to a VA to investigate + price the options.
    python3 scripts/gen_va_research_docx.py [output_path]
"""

import sys

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

GREEN = RGBColor(0x16, 0xA3, 0x4A)


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def num(doc, text):
    doc.add_paragraph(text, style="List Number")


def build(path):
    doc = Document()

    # Title
    t = doc.add_heading("PracticeRank — Digital PR / Earned Mentions", level=0)
    sub = doc.add_paragraph("Vendor Research Plan (VA Hand-off)")
    sub.runs[0].italic = True
    doc.add_paragraph("Goal: find the best + cheapest way to get our clients mentioned on "
                      "third-party websites that AI search engines (ChatGPT, Perplexity, "
                      "Google AI) cite. Research the options below, fill in the comparison "
                      "table, and recommend the top 1–2.")

    # 1. Context
    h(doc, "1. What we're trying to do (read first)")
    doc.add_paragraph(
        "When someone asks an AI assistant 'best gold buyer in Northern Virginia' (or "
        "'best dentist in Austin'), the AI doesn't only read our client's website — it "
        "pulls from third-party articles, news sites, and directories that MENTION the "
        "business. Research shows this off-site presence is the single biggest lever for "
        "AI visibility (about 3x more impactful than backlinks).")
    doc.add_paragraph(
        "So we need to get each client NAMED or QUOTED on other reputable websites. There "
        "are paid services that do this, and cheaper/free ways. Your job is to investigate "
        "each option, get real pricing, and tell us which to use.")
    doc.add_paragraph("Our clients are LOCAL service businesses: dental practices, law firms, "
                      "and local retail (e.g., a gold/silver buyer & jeweler). We charge them "
                      "~$3,000/month, so any cost here must stay small per client.")

    # 2. The models
    h(doc, "2. The three ways to buy this (plain English)")
    h(doc, "Model 1 — HARO / 'reporter looking for a source' services", level=2)
    doc.add_paragraph("Journalists post requests like 'I need a precious-metals expert for an "
                      "article on selling inherited jewelry.' Someone replies with a quote from "
                      "the business owner. If used, the business gets named in a real news "
                      "article — which AI then cites. HARO is free again (via Featured.com). A "
                      "paid 'service' monitors these daily and writes the responses for you.")
    h(doc, "Model 2 — Per-placement white-label digital PR ('buy a mention')", level=2)
    doc.add_paragraph("Providers do the outreach and land a mention/article on a real site, "
                      "billed PER PLACEMENT (market ~$200–350 each, range $49–800 by site "
                      "authority). 'White-label' = delivered unbranded so we resell it as "
                      "PracticeRank's work. No monthly retainer, no staff needed.")
    h(doc, "Model 3 — AEO/GEO done-for-you (à la carte)", level=2)
    doc.add_paragraph("Newer providers sell the off-site-mention work packaged specifically "
                      "for AI search — you buy just the 'off-site brand mentions' product per "
                      "client. Some have no minimums/retainers (quote required).")
    h(doc, "The cheap / DIY path (investigate this too)", level=2)
    doc.add_paragraph("HARO via Featured.com is FREE — we (or a VA) respond to journalist "
                      "requests; ~3–5 mentions/month with consistent pitching. Plus cheap press "
                      "wires (Send2Press ~$79, BrandPush ~$195) for guaranteed-but-lower-"
                      "authority syndication, and free direct pitching to local newspapers/blogs.")

    # 3. Options to investigate
    h(doc, "3. Specific options to investigate (visit + contact each)")

    h(doc, "Group A — HARO platforms & done-for-you HARO services", level=2)
    bullet(doc, "Featured.com (HARO) — featured.com — FREE platform. Confirm: free tier limits, how it works.")
    bullet(doc, "Qwoted — qwoted.com — ~$149/mo platform. Confirm current pricing + what paid unlocks.")
    bullet(doc, "FIND 3–5 'done-for-you HARO response services' (they monitor + write pitches for you). "
                "Search: 'HARO link building service', 'HARO response service for agencies', "
                "'managed HARO white label'. Get per-placement or monthly pricing for each.")

    h(doc, "Group B — Per-placement white-label digital PR", level=2)
    bullet(doc, "Reporter Outreach — reporteroutreach.com/white-label-link-building-services (3-month min; agency pricing on request).")
    bullet(doc, "RankZ — rankz.co (white-label link building).")
    bullet(doc, "SERPpro — serppro.ai/white-label-pr")
    bullet(doc, "Cedarwood Digital — cedarwood.digital/white-label-digital-pr")
    bullet(doc, "Also check this list for 3–4 more: orangeoutreach.com/best-white-label-link-building-agencies-usa")

    h(doc, "Group C — AEO/GEO done-for-you (built for AI search)", level=2)
    bullet(doc, "The AEO Collective — aeo-collective.com — sells off-site brand mentions à la carte. Get pricing for the OFF-SITE-MENTIONS product only.")
    bullet(doc, "E2M Solutions — e2msolutions.com/white-label-geo-services — no minimums/retainers; get a quote for off-site mentions standalone.")

    h(doc, "Group D — Cheap PR wires (guaranteed, low effort)", level=2)
    bullet(doc, "Send2Press — ~$79+ — confirm tiers + which outlets.")
    bullet(doc, "BrandPush — ~$195 — guaranteed 200–450 outlets, writing included.")
    bullet(doc, "Note quality: these are syndicated press releases (lower authority than editorial). Useful as a cheap baseline.")

    h(doc, "Group E — Free / DIY (for cost comparison)", level=2)
    bullet(doc, "HARO via Featured.com (free) — note effort: 20–40 pitches/mo for ~3–5 mentions.")
    bullet(doc, "Free PR sites: PR.com, PRLog, OpenPR (free, low authority).")
    bullet(doc, "Direct local-press pitching (local newspaper, business journal, niche blogs) — free, just outreach.")

    # 4. Questions
    h(doc, "4. Ask every paid vendor these 10 questions")
    for q in [
        "Exact wholesale/reseller price PER PLACEMENT, and any volume discounts?",
        "Is it per-placement, or a monthly minimum/retainer per client?",
        "Can you show 5–10 sample LIVE placements (so we can check the sites are real, indexed, and the kind AI cites — not link farms)?",
        "Fully white-label — nothing branded to you reaches our client?",
        "Do you draft + outreach + place end-to-end, or do we supply content?",
        "Is there an API or bulk portal to submit orders and get back the placement URL + date?",
        "Turnaround time per placement? Minimum term / cancellation policy?",
        "Can you land LOCAL/regional + NICHE outlets (e.g., precious metals, dental, legal)?",
        "Do you guarantee placement, or best-effort? Refund/replace policy if not placed?",
        "Any tactics that could risk a Google penalty for our clients (PBNs, paid-link footprints)?",
    ]:
        num(doc, q)

    # 5. Comparison table
    h(doc, "5. Fill in this comparison table")
    cols = ["Vendor", "Model (1-4)", "$ / placement", "Per-placement? (Y/N)",
            "White-label? (Y/N)", "Turnkey? (Y/N)", "API/portal? (Y/N)", "Min term",
            "Sample sites quality (1-5)", "Turnaround", "Score (1-10)"]
    rows = ["Featured.com (HARO)", "Qwoted", "HARO service #1", "HARO service #2",
            "Reporter Outreach", "RankZ", "SERPpro", "Cedarwood",
            "The AEO Collective", "E2M Solutions", "Send2Press", "BrandPush"]
    table = doc.add_table(rows=1 + len(rows), cols=len(cols))
    table.style = "Light Grid Accent 1"
    for j, c in enumerate(cols):
        cell = table.rows[0].cells[j]
        cell.text = c
        for r in cell.paragraphs[0].runs:
            r.bold = True
    for i, name in enumerate(rows, start=1):
        table.rows[i].cells[0].text = name

    # 6. What we want
    h(doc, "6. What we're looking for (our criteria)")
    bullet(doc, "PER-PLACEMENT or à la carte — NOT a monthly retainer per client.")
    bullet(doc, "Cheap: ideally under ~$200/placement; target blended cost under ~$200–350 per client per month.")
    bullet(doc, "White-label (we resell as PracticeRank).")
    bullet(doc, "Turnkey — they do the work; no staff needed on our side.")
    bullet(doc, "Real editorial / news / indexed sites that AI actually cites — NOT spammy link farms.")
    bullet(doc, "Bonus: an API/portal so we can automate orders, and they return the placement URL.")

    # 7. Deliverable
    h(doc, "7. What to hand back (deliverable)")
    num(doc, "The comparison table above, filled in for every vendor you reached.")
    num(doc, "Real pricing for each (per placement + any minimums). Note who wouldn't share pricing.")
    num(doc, "3–5 done-for-you HARO services you found, with pricing.")
    num(doc, "For the top 2–3, attach the sample live placements they sent (so we can judge quality).")
    num(doc, "A short recommendation: which 1–2 you'd pick and why (best quality-per-dollar), plus the cheapest viable option.")

    # 8. Notes
    h(doc, "8. Helpful notes")
    bullet(doc, "Timeline context: HARO mentions usually publish within 1–2 weeks; paid placements take ~2–4 weeks each.")
    bullet(doc, "Quality check: always ask for sample live URLs and open them — confirm they're real sites that show up in Google, not auto-generated 'news' pages.")
    bullet(doc, "Avoid: full digital-PR retainers ($2,000–10,000/month) — too expensive for our model.")
    bullet(doc, "Ignore: 'AI tracking' tools (LLM Pulse, AI Rank Lab) — we already built our own tracking.")

    doc.save(path)
    print(f"Wrote {path}")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "docs/Lever-A-VA-Research-Plan.docx"
    build(out)
