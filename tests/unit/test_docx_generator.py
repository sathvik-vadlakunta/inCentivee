"""Tests for DOCX content deliverable generation."""

import pytest
from io import BytesIO
from docx import Document

from geo_agent.docx_generator import ContentDocxGenerator


@pytest.fixture
def customer():
    return {"id": "test-dental", "name": "Test Dental", "domain": "testdental.com", "platform": "webflow"}


@pytest.fixture
def generator(customer):
    return ContentDocxGenerator(customer)


@pytest.fixture
def blog_rec():
    return {
        "id": "rec-1",
        "customer_id": "test-dental",
        "rec_type": "blog_post",
        "title": "5 Tips for Better Oral Hygiene",
        "description": "Learn the top oral hygiene tips from our experts.",
        "html_snippet": "<h2>Tip 1: Brush Twice Daily</h2><p>Brushing <strong>twice a day</strong> is essential for <em>maintaining</em> healthy teeth.</p><ul><li>Use a soft-bristled brush</li><li>Brush for 2 minutes</li></ul>",
        "priority": 1,
        "category": "general",
        "status": "approved",
    }


@pytest.fixture
def faq_rec():
    return {
        "id": "rec-2",
        "customer_id": "test-dental",
        "rec_type": "faq_update",
        "title": "FAQ: Dental Implants",
        "description": "Common questions about dental implants.",
        "html_snippet": "<p><strong>Q: How long do dental implants last?</strong></p><p>A: With proper care, dental implants can last a lifetime.</p>",
        "target_page": "/services/implants",
        "priority": 2,
        "category": "general",
        "status": "approved",
    }


@pytest.fixture
def stat_rec():
    return {
        "id": "rec-3",
        "customer_id": "test-dental",
        "rec_type": "stat_injection",
        "title": "Add Success Rate Stat",
        "description": "Insert implant success rate statistic on services page.",
        "html_snippet": "<p>Dental implants have a <strong>98% success rate</strong> according to the ADA.</p>",
        "target_page": "/services/implants",
        "priority": 2,
        "category": "general",
        "status": "approved",
    }


@pytest.fixture
def expert_quote_rec():
    return {
        "id": "rec-4",
        "customer_id": "test-dental",
        "rec_type": "expert_quote",
        "title": "Dr. Smith Quote on Implants",
        "description": "Expert quote for implants page.",
        "html_snippet": '<blockquote>"Modern dental implants are the gold standard for tooth replacement." — Dr. Smith, DDS</blockquote>',
        "target_page": "/services/implants",
        "priority": 3,
        "category": "general",
        "status": "approved",
    }


class TestSingleGeneration:
    def test_blog_post_generates_valid_docx(self, generator, blog_rec):
        buf = generator.generate_single(blog_rec)
        assert isinstance(buf, BytesIO)
        doc = Document(buf)
        # Should have content
        assert len(doc.paragraphs) > 0
        # Title should be in headings
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "5 Tips for Better Oral Hygiene" in full_text

    def test_faq_generates_valid_docx(self, generator, faq_rec):
        buf = generator.generate_single(faq_rec)
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "FAQ: Dental Implants" in full_text
        assert "/services/implants" in full_text

    def test_stat_injection_generates_valid_docx(self, generator, stat_rec):
        buf = generator.generate_single(stat_rec)
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Add Success Rate Stat" in full_text

    def test_expert_quote_generates_valid_docx(self, generator, expert_quote_rec):
        buf = generator.generate_single(expert_quote_rec)
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Dr. Smith" in full_text

    def test_new_page_generates_valid_docx(self, generator):
        rec = {
            "id": "rec-5",
            "rec_type": "new_page",
            "title": "Emergency Dental Services",
            "description": "24/7 emergency dental care in Austin.",
            "html_snippet": "<p>We offer emergency dental care.</p>",
            "priority": 1,
            "status": "approved",
        }
        buf = generator.generate_single(rec)
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Emergency Dental Services" in full_text

    def test_freshness_update_generates_valid_docx(self, generator):
        rec = {
            "id": "rec-6",
            "rec_type": "freshness_update",
            "title": "Update Insurance Page",
            "description": "Insurance list is outdated.",
            "html_snippet": "<p>We now accept Delta Dental and Cigna.</p>",
            "target_page": "/insurance",
            "priority": 2,
            "status": "approved",
        }
        buf = generator.generate_single(rec)
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Update Insurance Page" in full_text


class TestBatchGeneration:
    def test_batch_generates_valid_docx(self, generator, blog_rec, faq_rec, stat_rec):
        buf = generator.generate_batch([blog_rec, faq_rec, stat_rec])
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "5 Tips for Better Oral Hygiene" in full_text
        assert "FAQ: Dental Implants" in full_text
        assert "Add Success Rate Stat" in full_text

    def test_empty_batch(self, generator):
        buf = generator.generate_batch([])
        doc = Document(buf)
        # Should just have header
        assert len(doc.paragraphs) >= 2


class TestHtmlToDOCX:
    def test_handles_bold_text(self, generator):
        doc = Document()
        generator._html_to_docx(doc, "<p>This is <strong>bold</strong> text.</p>")
        assert len(doc.paragraphs) > 0

    def test_handles_headings(self, generator):
        doc = Document()
        generator._html_to_docx(doc, "<h2>Main Heading</h2><h3>Sub Heading</h3><p>Content here.</p>")
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Main Heading" in full_text
        assert "Sub Heading" in full_text

    def test_handles_lists(self, generator):
        doc = Document()
        generator._html_to_docx(doc, "<ul><li>Item 1</li><li>Item 2</li></ul>")
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Item 1" in full_text
        assert "Item 2" in full_text

    def test_handles_blockquote(self, generator):
        doc = Document()
        generator._html_to_docx(doc, '<blockquote>"A wise quote" — Author</blockquote>')
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "A wise quote" in full_text

    def test_handles_empty_html(self, generator):
        doc = Document()
        generator._html_to_docx(doc, "")
        # Should not crash

    def test_handles_plain_text(self, generator):
        doc = Document()
        generator._html_to_docx(doc, "Just some plain text without tags")
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Just some plain text" in full_text


class TestFaqSchemaOrg:
    """Test FAQ rendering with Schema.org FAQPage markup (nested divs)."""

    def test_section_title_not_rendered_as_question(self, generator):
        """H2 section titles with sub-headings should render as headings, not Q&A."""
        html = (
            '<div itemscope itemtype="https://schema.org/FAQPage">'
            '<h2>Frequently Asked Questions About Implants</h2>'
            '<div itemscope itemprop="mainEntity" itemtype="https://schema.org/Question">'
            '<h3 itemprop="name">How long do implants last?</h3>'
            '<div itemscope itemprop="acceptedAnswer" itemtype="https://schema.org/Answer">'
            '<div itemprop="text"><p>With proper care, implants can last a lifetime.</p></div>'
            '</div></div>'
            '<div itemscope itemprop="mainEntity" itemtype="https://schema.org/Question">'
            '<h3 itemprop="name">Are implants painful?</h3>'
            '<div itemscope itemprop="acceptedAnswer" itemtype="https://schema.org/Answer">'
            '<div itemprop="text"><p>Most patients report minimal discomfort.</p></div>'
            '</div></div></div>'
        )
        rec = {
            "id": "faq-schema", "rec_type": "faq_update",
            "title": "Implant FAQ", "html_snippet": html,
            "priority": 1, "status": "approved",
        }
        buf = generator.generate_single(rec)
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        # Section title should NOT have "Q:" prefix
        assert "Q: Frequently Asked Questions" not in full_text
        # But actual questions should
        assert "Q: How long do implants last?" in full_text
        assert "Q: Are implants painful?" in full_text
        # Answers should be present
        assert "implants can last a lifetime" in full_text
        assert "minimal discomfort" in full_text

    def test_heading_without_answer_rendered_as_heading(self, generator):
        """Headings with no answer content should be section headings, not Q&A."""
        html = '<h2>Our FAQ Section</h2><h3>What do you offer?</h3><p>We offer everything.</p>'
        rec = {
            "id": "faq-no-answer", "rec_type": "faq_update",
            "title": "FAQ No Answer", "html_snippet": html,
            "priority": 2, "status": "approved",
        }
        buf = generator.generate_single(rec)
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Q: Our FAQ Section" not in full_text
        assert "Q: What do you offer?" in full_text
        assert "We offer everything" in full_text


class TestBranding:
    def test_header_includes_brand_and_customer(self, generator):
        doc = Document()
        generator._add_header(doc)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "PracticeRank" in full_text
        assert "Test Dental" in full_text
