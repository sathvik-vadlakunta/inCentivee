"""Generate DOCX content deliverables from content recommendations.

Works for all platforms — downloadable for VA/client review before publishing.
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class ContentDocxGenerator:
    """Generate branded DOCX files from content recommendations."""

    def __init__(self, customer: dict, brand_name: str = "PracticeRank"):
        self.customer = customer
        self.brand_name = brand_name

    def generate_single(self, rec: dict) -> BytesIO:
        """Generate a DOCX for a single recommendation."""
        doc = Document()
        self._add_header(doc)
        self._render_rec(doc, rec)
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    def generate_batch(self, recs: list[dict]) -> BytesIO:
        """Generate a single DOCX containing all recommendations with page breaks."""
        doc = Document()
        self._add_header(doc)
        for i, rec in enumerate(recs):
            if i > 0:
                doc.add_page_break()
            self._render_rec(doc, rec)
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    def _render_rec(self, doc: Document, rec: dict):
        """Dispatch to type-specific renderer."""
        rec_type = rec.get("rec_type", "blog_post")
        renderer = {
            "blog_post": self._render_blog_post,
            "faq_update": self._render_faq_update,
            "stat_injection": self._render_stat_injection,
            "freshness_update": self._render_freshness_update,
            "expert_quote": self._render_expert_quote,
            "new_page": self._render_new_page,
        }.get(rec_type, self._render_blog_post)
        renderer(doc, rec)

    def _render_blog_post(self, doc: Document, rec: dict):
        """Full article with title, meta description, body."""
        self._add_type_badge(doc, "Blog Post", rec.get("priority", 3))
        doc.add_heading(rec.get("title", "Untitled"), level=1)
        if rec.get("description"):
            p = doc.add_paragraph()
            run = p.add_run(f"Meta Description: {rec['description']}")
            run.font.size = Pt(10)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        if rec.get("html_snippet"):
            self._html_to_docx(doc, rec["html_snippet"])

    def _render_faq_update(self, doc: Document, rec: dict):
        """Q&A pairs with clear question headers and answer blocks."""
        self._add_type_badge(doc, "FAQ Update", rec.get("priority", 3))
        doc.add_heading(rec.get("title", "FAQ Update"), level=1)
        if rec.get("target_page"):
            doc.add_paragraph(f"Target Page: {rec['target_page']}")
        if rec.get("html_snippet"):
            self._faq_html_to_docx(doc, rec["html_snippet"])

    def _render_stat_injection(self, doc: Document, rec: dict):
        """Target page + what to change."""
        self._add_type_badge(doc, "Stat Injection", rec.get("priority", 3))
        doc.add_heading(rec.get("title", "Stat Injection"), level=1)
        if rec.get("target_page"):
            doc.add_paragraph(f"Target Page: {rec['target_page']}")
        if rec.get("description"):
            doc.add_paragraph(rec["description"])
        if rec.get("html_snippet"):
            doc.add_heading("Content to Insert:", level=2)
            self._html_to_docx(doc, rec["html_snippet"])

    def _render_freshness_update(self, doc: Document, rec: dict):
        """Instruction card: page, what's stale, replacement."""
        self._add_type_badge(doc, "Freshness Update", rec.get("priority", 3))
        doc.add_heading(rec.get("title", "Freshness Update"), level=1)
        if rec.get("target_page"):
            doc.add_paragraph(f"Page to Update: {rec['target_page']}")
        if rec.get("description"):
            doc.add_paragraph(rec["description"])
        if rec.get("html_snippet"):
            doc.add_heading("Updated Content:", level=2)
            self._html_to_docx(doc, rec["html_snippet"])

    def _render_expert_quote(self, doc: Document, rec: dict):
        """Quote + attribution + insertion location."""
        self._add_type_badge(doc, "Expert Quote", rec.get("priority", 3))
        doc.add_heading(rec.get("title", "Expert Quote"), level=1)
        if rec.get("target_page"):
            doc.add_paragraph(f"Insert on Page: {rec['target_page']}")
        if rec.get("html_snippet"):
            self._html_to_docx(doc, rec["html_snippet"])

    def _render_new_page(self, doc: Document, rec: dict):
        """Same as blog_post, labeled as standalone page."""
        self._add_type_badge(doc, "New Page", rec.get("priority", 3))
        doc.add_heading(rec.get("title", "New Page"), level=1)
        if rec.get("description"):
            p = doc.add_paragraph()
            run = p.add_run(f"Meta Description: {rec['description']}")
            run.font.size = Pt(10)
            run.font.italic = True
        if rec.get("html_snippet"):
            self._html_to_docx(doc, rec["html_snippet"])

    def _add_header(self, doc: Document):
        """Add PracticeRank branding header."""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(self.brand_name)
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

        p2 = doc.add_paragraph()
        p2.add_run(f"Customer: {self.customer.get('name', 'Unknown')}").bold = True
        p2.add_run(f"  |  Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d')}")
        doc.add_paragraph()  # spacer

    def _add_type_badge(self, doc: Document, type_name: str, priority: int):
        """Add a type/priority line."""
        p = doc.add_paragraph()
        run = p.add_run(f"[{type_name}]")
        run.bold = True
        run.font.size = Pt(9)
        priority_colors = {1: RGBColor(0xDC, 0x26, 0x26), 2: RGBColor(0xD9, 0x77, 0x06), 3: RGBColor(0x6B, 0x72, 0x80)}
        run.font.color.rgb = priority_colors.get(priority, RGBColor(0x6B, 0x72, 0x80))
        p.add_run(f"  Priority {priority}").font.size = Pt(9)

    def _html_to_docx(self, doc: Document, html: str):
        """Convert well-structured HTML to DOCX paragraphs in document order."""
        if not html:
            return

        html = html.strip()

        # Tokenize into block-level elements in document order
        block_pattern = re.compile(
            r'(<h([2-4])[^>]*>.*?</h\2>)'       # headings
            r'|(<blockquote[^>]*>.*?</blockquote>)'  # blockquotes
            r'|(<[ou]l[^>]*>.*?</[ou]l>)'         # lists
            r'|(<p[^>]*>.*?</p>)',                 # paragraphs
            re.DOTALL
        )

        last_end = 0
        has_blocks = False

        for match in block_pattern.finditer(html):
            has_blocks = True

            # Handle any text between blocks
            gap = html[last_end:match.start()].strip()
            gap_text = self._strip_tags(gap).strip()
            if gap_text:
                doc.add_paragraph(gap_text)

            last_end = match.end()
            tag_html = match.group(0)

            # Heading
            h_match = re.match(r'<h([2-4])[^>]*>(.*?)</h\1>', tag_html, re.DOTALL)
            if h_match:
                level = int(h_match.group(1))
                text = self._strip_tags(h_match.group(2))
                doc.add_heading(text, level=level)
                continue

            # Blockquote
            bq_match = re.match(r'<blockquote[^>]*>(.*?)</blockquote>', tag_html, re.DOTALL)
            if bq_match:
                text = self._strip_tags(bq_match.group(1))
                doc.add_paragraph(text, style="Quote")
                continue

            # List
            list_match = re.match(r'<[ou]l[^>]*>(.*?)</[ou]l>', tag_html, re.DOTALL)
            if list_match:
                items = re.findall(r'<li[^>]*>(.*?)</li>', list_match.group(1), re.DOTALL)
                for item in items:
                    text = self._strip_tags(item)
                    doc.add_paragraph(text, style="List Bullet")
                continue

            # Paragraph
            p_match = re.match(r'<p[^>]*>(.*?)</p>', tag_html, re.DOTALL)
            if p_match:
                p = doc.add_paragraph()
                self._add_inline_runs(p, p_match.group(1))
                continue

        # Handle trailing text after last block
        if has_blocks:
            trailing = html[last_end:].strip()
            trailing_text = self._strip_tags(trailing).strip()
            if trailing_text:
                doc.add_paragraph(trailing_text)
        else:
            # No block tags found — treat as plain text
            text = self._strip_tags(html).strip()
            if text:
                for line in text.split('\n'):
                    line = line.strip()
                    if line:
                        doc.add_paragraph(line)

    def _faq_html_to_docx(self, doc: Document, html: str):
        """Render FAQ HTML with clear Q/A block structure.

        Detects question headings (h2-h4) and treats subsequent paragraphs as answers,
        adding visual separation between each Q&A pair.
        Handles Schema.org FAQPage markup with nested div structure.
        """
        if not html:
            return

        html = html.strip()

        # Try to extract Q&A pairs from heading+paragraph structure
        # Pattern: <h2-4>Question</h2-4> followed by one or more <p>Answer</p>
        qa_pattern = re.compile(
            r'<h([2-4])[^>]*>(.*?)</h\1>(.*?)(?=<h[2-4]|$)',
            re.DOTALL
        )
        matches = list(qa_pattern.finditer(html))

        if matches:
            # Extract any content before the first Q (like a section title)
            preamble = html[:matches[0].start()].strip()
            if preamble:
                self._html_to_docx(doc, preamble)

            for match in matches:
                heading_level = int(match.group(1))
                question = self._strip_tags(match.group(2)).strip()
                answer_html = match.group(3).strip()

                # Extract answer <p> tags (may be nested inside schema.org divs)
                answer_parts = re.findall(r'<p[^>]*>(.*?)</p>', answer_html, re.DOTALL)

                # If this is an h2 with no direct <p> answers but contains h3s,
                # it's a section title — render as heading, not Q&A
                has_sub_headings = re.search(r'<h[3-4]', answer_html)
                if heading_level == 2 and not answer_parts and has_sub_headings:
                    doc.add_heading(question, level=2)
                    continue

                # If no <p> tags, try stripping all tags for plain text answer
                if not answer_parts:
                    plain = self._strip_tags(answer_html).strip()
                    if not plain:
                        # No answer at all — render as section heading, not Q
                        doc.add_heading(question, level=heading_level)
                        continue

                # Question as bold heading with "Q:" prefix
                q_para = doc.add_paragraph()
                q_para.space_before = Pt(12)
                q_run = q_para.add_run(f"Q: {question}")
                q_run.bold = True
                q_run.font.size = Pt(11)
                q_run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

                # Answer paragraphs
                if answer_parts:
                    for part in answer_parts:
                        a_para = doc.add_paragraph()
                        a_run = a_para.add_run("A: ")
                        a_run.bold = True
                        a_run.font.size = Pt(10)
                        self._add_inline_runs(a_para, part)
                else:
                    plain = self._strip_tags(answer_html).strip()
                    if plain:
                        a_para = doc.add_paragraph()
                        a_run = a_para.add_run("A: ")
                        a_run.bold = True
                        a_run.font.size = Pt(10)
                        a_para.add_run(plain)

                # Add spacing after each Q&A pair
                spacer = doc.add_paragraph()
                spacer.space_after = Pt(4)
        else:
            # Fallback: no heading structure, use regular rendering
            self._html_to_docx(doc, html)

    def _add_inline_runs(self, paragraph, html: str):
        """Add runs with bold/italic formatting from inline HTML.

        Preserves whitespace around inline elements so "The <a>link</a> text"
        renders as "The link text" not "Thelinktext".
        """
        # Simple pattern: split on <strong>, <em>, <a>, <cite> tags
        parts = re.split(r'(<(?:strong|em|b|i|a|cite)[^>]*>.*?</(?:strong|em|b|i|a|cite)>)', html, flags=re.DOTALL)
        for part in parts:
            if not part:
                continue
            strong_match = re.match(r'<(?:strong|b)[^>]*>(.*?)</(?:strong|b)>', part, re.DOTALL)
            em_match = re.match(r'<(?:em|i)[^>]*>(.*?)</(?:em|i)>', part, re.DOTALL)
            a_match = re.match(r'<a[^>]*>(.*?)</a>', part, re.DOTALL)
            cite_match = re.match(r'<cite[^>]*>(.*?)</cite>', part, re.DOTALL)
            if strong_match:
                run = paragraph.add_run(self._strip_tags(strong_match.group(1)))
                run.bold = True
            elif em_match:
                run = paragraph.add_run(self._strip_tags(em_match.group(1)))
                run.italic = True
            elif a_match:
                run = paragraph.add_run(self._strip_tags(a_match.group(1)))
                run.underline = True
                run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            elif cite_match:
                run = paragraph.add_run(self._strip_tags(cite_match.group(1)))
                run.italic = True
                run.font.size = Pt(9)
            else:
                # Preserve whitespace — don't strip so spaces around
                # inline elements are kept (e.g. "The <a>link</a> text")
                text = self._strip_tags(part, strip_whitespace=False)
                if text:
                    paragraph.add_run(text)

    @staticmethod
    def _strip_tags(html: str, strip_whitespace: bool = True) -> str:
        """Remove all HTML tags, decode common entities."""
        text = re.sub(r'<[^>]+>', '', html)
        text = text.replace('&amp;', '&').replace('&lt;', '<').replace('&gt;', '>')
        text = text.replace('&nbsp;', ' ').replace('&#39;', "'").replace('&quot;', '"')
        return text.strip() if strip_whitespace else text
