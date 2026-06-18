#!/usr/bin/env python3
"""Generate the Local AI/Search Playbook (.docx) from the verified deep-research.

Three topics: Reddit-for-local, high-value local links, Google Maps/3-pack.
Marks each item VERIFIED / CONVENTIONAL (unverified) / DEBUNKED so staff don't
work against refuted claims.
    python3 scripts/gen_local_playbook_docx.py [output_path]
"""

import sys

from docx import Document


def h(doc, t, level=1):
    doc.add_heading(t, level=level)


def b(doc, t):
    doc.add_paragraph(t, style="List Bullet")


def n(doc, t):
    doc.add_paragraph(t, style="List Number")


def build(path):
    doc = Document()
    doc.add_heading("PracticeRank — Local AI & Search Playbook", level=0)
    p = doc.add_paragraph("Evidence-backed (2025–2026) playbook for LOCAL clients — dental, "
                          "legal, local retail. Each item is tagged so staff know what's proven "
                          "vs conventional wisdom vs debunked.")
    p.runs[0].italic = True
    doc.add_paragraph("Legend:  [VERIFIED] = survived adversarial fact-check (cited).  "
                      "[CONVENTIONAL] = reasonable industry practice but NOT proven by the "
                      "evidence we gathered — use, but don't claim as fact.  [DEBUNKED] = "
                      "failed verification — do NOT staff against it.")

    # Cross-cutting reality
    h(doc, "0. The big picture (read first)")
    b(doc, "[VERIFIED] AI is now a mainstream but SELECTIVE local-discovery channel: ~45% of "
           "consumers used AI for local recommendations in early 2026 (up from 6%), led by "
           "ChatGPT (31%) and Google AI Mode (23%). BUT AI 'local packs' surface only ~32% as "
           "many businesses as Google's 3-pack, and fewer than half of Google local leaders "
           "appear in AI results. (BrightLocal LCRS 2026; Sterling Sky; SOCi.)")
    b(doc, "[VERIFIED] Implication: ranking #1 in Google Maps does NOT mean you'll be the AI "
           "recommendation. To get surfaced by AI for local queries you must be present across "
           "MANY third-party platforms — directories, Yelp (cited in ~33% of AI searches), "
           "reviews, social, and your own site — not just Google.")
    b(doc, "[CAVEAT] These figures move fast: Google AI Mode's self-citation of GBP fell from "
           "~98% (mid-2025) to ~36% (Feb 2026). Re-verify quarterly.")

    # Topic 3 first (strongest evidence)
    h(doc, "1. Google Maps / Local 3-Pack ranking (strongest evidence)")
    h(doc, "What actually moves the needle", level=2)
    b(doc, "[VERIFIED] Google Business Profile is the foundation. Set the right PRIMARY category; "
           "add relevant ADDITIONAL categories — profiles with ~4 additional categories averaged "
           "map position 5.9 vs 7.6 with none (correlation, not proven causation).")
    b(doc, "[VERIFIED] Reviews are the #2 controllable factor (grew 16%→20% of local-pack weight). "
           "What matters: number of NATIVE Google reviews WITH TEXT, recency, and steady growth.")
    b(doc, "[VERIFIED] There's a threshold around 10 reviews: going 9→10 gave a noticeable bump; "
           "10→11 did not. Get every client to ~10+ Google reviews, then keep a steady drip "
           "(diminishing returns on raw count beyond ~10).")
    b(doc, "[VERIFIED] Review recency is also a conversion signal: 74% of consumers want reviews "
           "from the last 3 months; 32% from the last 2 weeks.")
    b(doc, "[VERIFIED] Citation / NAP consistency matters: a dentist with 11 of 13 citations "
           "containing errors scored 28/100; 62% of consumers avoid a business with incorrect "
           "info online. Keep Name/Address/Phone identical everywhere.")
    b(doc, "[VERIFIED] Proximity to the searcher is a heavy factor — uncontrollable, but it's why "
           "service-area pages + being correctly placed/categorized matters.")
    h(doc, "What NOT to bill as ranking levers (no proven ranking impact)", level=2)
    b(doc, "[DEBUNKED] Geo-tagged photos, keywords stuffed into review RESPONSES, keywords in the "
           "GBP description, sheer QUANTITY of Google Posts, and QUANTITY of Q&A entries — none "
           "showed a ranking effect. (They can still help conversion/trust — sell them as that, "
           "not as ranking factors.)")
    b(doc, "[DEBUNKED] 'Review velocity beats total count' and 'rankings drop if you stop getting "
           "reviews for ~3 weeks' — both failed verification. Don't promise these.")
    b(doc, "[DEBUNKED] 'Reviews with keywords beat star-only ratings' — not supported.")
    h(doc, "Staff playbook (Maps)", level=2)
    n(doc, "Verify + fully complete GBP; set primary + 3-4 accurate additional categories.")
    n(doc, "Fill services, hours, attributes, photos (for conversion, not ranking).")
    n(doc, "Drive to ~10+ native Google reviews WITH text, then a steady monthly drip; respond to all.")
    n(doc, "Lock NAP and fix citation errors across the core + vertical directories.")
    n(doc, "Build city/service-area pages so relevance + proximity signals line up.")

    # Topic 2
    h(doc, "2. High-value local links / mentions (quality over volume)")
    doc.add_paragraph("IMPORTANT: the specific local-link tactics below did NOT survive "
                      "fact-checking with citable evidence in this research batch. Treat them as "
                      "CONVENTIONAL best practice, not proven — and prioritize the VERIFIED "
                      "'broad third-party presence + citation consistency' finding above.")
    b(doc, "[VERIFIED, adjacent] The proven 'link-like' lever for local is broad, consistent "
           "presence across directories + Yelp + reviews + your site — that's what AI cites for "
           "local. Nail citations/NAP first.")
    b(doc, "[CONVENTIONAL] Local news / press features (pitch a local-angle story).")
    b(doc, "[CONVENTIONAL] Local .edu / .gov pages, chambers of commerce, local business associations.")
    b(doc, "[CONVENTIONAL] Local sponsorships, events, and scholarships (often earn a local-org link).")
    b(doc, "[CONVENTIONAL] Supplier / partner / vendor pages that list the business.")
    b(doc, "[CONVENTIONAL] '​Best {service} in {city}' roundup inclusion; local podcasts.")
    b(doc, "Principle: a few HIGH-relevance LOCAL links beat volume — but we can't cite hard "
           "numbers, so don't quote stats; pitch it as quality-over-quantity hygiene.")

    # Topic 1
    h(doc, "3. Reddit for local businesses")
    b(doc, "[VERIFIED] Reddit is cited heavily by AI for GENERAL/commercial queries (~24% of "
           "Perplexity citations; 44% of social citations in Google AI Overviews).")
    b(doc, "[DEBUNKED] 'Reddit/Quora mentions give ~4x higher AI-citation odds' — failed verification.")
    b(doc, "[UNPROVEN for local] Whether Reddit helps for LOCAL '{service} in {city}' queries was "
           "NOT shown by the evidence. The heavy-citation data is national/commercial, not local.")
    b(doc, "Recommendation: for a single-location client, Reddit is LOW priority. If anything, do "
           "light, genuine participation in the city subreddit (answer local questions honestly, "
           "no spam — most local subs ban self-promotion). Do not build a paid Reddit motion or "
           "promise AI-citation lift from it until proven.")

    # Caveats + sources
    h(doc, "4. Caveats")
    b(doc, "AI-local stats move fast (GBP self-citation 98%→36% in months) — re-verify quarterly.")
    b(doc, "The 45% AI-usage figure is a single US survey (n≈1,002, self-report); treat the YoY jump magnitude loosely.")
    b(doc, "Sterling Sky / SOCi figures are from multi-location brands — directionally true for solo SMBs, exact %s less certain.")
    b(doc, "Review-threshold test used only 3 businesses/step (small, correlational).")

    h(doc, "5. Sources")
    for s in [
        "BrightLocal — LCRS AI Trust 2026: brightlocal.com/research/lcrs-ai-trust/",
        "BrightLocal — AI search using listings/sources: brightlocal.com/blog/ai-search-using-listings-sources/",
        "BrightLocal — Local ranking factors / Google local algorithm: brightlocal.com/learn/google-local-algorithm-and-ranking-factors/",
        "BrightLocal — Dentist local rankings investigation (categories, citations): brightlocal.com/research/local-rankings-investigation-dentist/",
        "BrightLocal — Local Consumer Review Survey: brightlocal.com/research/local-consumer-review-survey/",
        "Sterling Sky — State of Local SEO 2026: sterlingsky.ca/the-state-of-local-seo-in-2026/",
        "Sterling Sky — Number of reviews impact on ranking: sterlingsky.ca/number-of-reviews-impact-ranking/",
        "CMSWire — Reddit's rise in AI citations: cmswire.com/digital-marketing/reddits-rise-in-ai-citations...",
    ]:
        b(doc, s)

    doc.save(path)
    print(f"Wrote {path}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "docs/Local-AI-Search-Playbook.docx")
