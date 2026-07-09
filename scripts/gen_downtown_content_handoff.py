#!/usr/bin/env python3
"""Build Downtown Dental's content-to-deploy handoff DOCX from an exported JSON of
approved content recommendations. Downtown self-publishes (we provide content,
their dev applies it), so this packages every approved snippet by target page.

    python scripts/gen_downtown_content_handoff.py <content.json> [out.docx]
"""
import json
import sys
from collections import defaultdict

from docx import Document
from docx.shared import Pt, RGBColor

DARK = RGBColor(0x14, 0x2A, 0x3D)
GREY = RGBColor(0x55, 0x5B, 0x66)
ACCENT = RGBColor(0x1F, 0x6F, 0xB2)

TYPE_LABEL = {
    "faq_update": "FAQ section (accordion Q&A + schema)",
    "blog_post": "Blog post",
    "stat_injection": "Statistic callout",
    "freshness_update": "Freshness update",
    "new_page": "New page",
}


def _mono(doc, text):
    for line in (text or "").splitlines() or [""]:
        p = doc.add_paragraph()
        r = p.add_run(line or " ")
        r.font.name = "Consolas"
        r.font.size = Pt(7.5)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(6)


def build(items, out_path):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    h = doc.add_heading(level=0)
    h.add_run("Downtown Dental — Content to Deploy (July 2026)").font.color.rgb = DARK
    sub = doc.add_paragraph()
    s = sub.add_run(
        f"{len(items)} approved content items, prepared by PracticeRank for Downtown Dental's "
        "web developer to publish. Grouped by the page each item belongs on."
    )
    s.font.color.rgb = GREY
    s.font.size = Pt(10)

    note = doc.add_paragraph()
    note.add_run("How to publish").bold = True
    for n in [
        "Paste each snippet onto its named target page (headings below are the page URLs).",
        "Publish the HTML EXACTLY as given. For FAQ items, keep every <details>/<summary> "
        "block separate — one collapsible question each — and keep the schema attributes "
        "(itemprop=\"mainEntity\"/Question/acceptedAnswer). Do NOT flatten them into one block: "
        "that drops the per-question schema AI engines cite (this happened to the Dr. Zhivago FAQ).",
        "Blog posts go up as new posts under /blog/.",
        "Statistic callouts insert into the existing page copy where noted; freshness updates "
        "refresh the existing page text.",
    ]:
        doc.add_paragraph(n, style="List Bullet")
    doc.add_paragraph()

    # Group by target page; blog last.
    groups = defaultdict(list)
    for it in items:
        groups[it.get("target_page") or "(unspecified)"].append(it)

    def _key(pg):
        return (1 if "/blog" in pg else 0, pg)

    for pg in sorted(groups, key=_key):
        ph = doc.add_heading(level=2)
        ph.add_run(pg).font.color.rgb = ACCENT
        for it in groups[pg]:
            th = doc.add_heading(it.get("title") or "(untitled)", level=3)
            meta = doc.add_paragraph()
            bits = [TYPE_LABEL.get(it.get("rec_type"), it.get("rec_type") or "")]
            if it.get("intent_tier"):
                bits.append(f"intent: {it['intent_tier']}")
            mr = meta.add_run("  ·  ".join(bits))
            mr.font.size = Pt(9)
            mr.font.color.rgb = GREY
            if it.get("meta_description"):
                md = doc.add_paragraph()
                mdr = md.add_run("Meta: " + it["meta_description"])
                mdr.font.size = Pt(9)
                mdr.italic = True
                mdr.font.color.rgb = GREY
            _mono(doc, it.get("html_snippet") or "")
            doc.add_paragraph()

    doc.save(out_path)
    print(f"wrote {out_path} ({len(items)} items across {len(groups)} pages)")


if __name__ == "__main__":
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "docs/Downtown-Dental-Content-to-Deploy-2026-07.docx"
    with open(src) as f:
        build(json.load(f), out)
