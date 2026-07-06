#!/usr/bin/env python3
"""Generate a PracticeRank-branded expense / reimbursement report as a polished .docx.

This is the reusable TEMPLATE for monthly out-of-pocket reimbursement submissions.

Two modes:

  # 1. Reproduce a fixed report from the embedded data (default = May-June 2026)
  python3 scripts/gen_expense_report_docx.py [out.docx]

  # 2. Regenerate live from the expense ledger for any months (future months):
  python3 scripts/gen_expense_report_docx.py --db data/practicerank.db \
          --months 2026-07,2026-08 --period "July 1 - Aug 31, 2026" [out.docx]

The ledger lives on the droplet at /app/data/practicerank.db; copy it down or run
this inside the container. Basis is inferred: prepaid-credit vendors (Perplexity,
xAI/Grok) -> "Prepaid", the recurring flag -> "Recurring", FATJOE offsite orders ->
"COGS", everything else -> "Usage".
"""

from __future__ import annotations

import argparse
import datetime as _dt
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

# ── PracticeRank brand ──────────────────────────────────────────────────────
GREEN = RGBColor(0x16, 0xA3, 0x4A)   # --accent
GREEN_DK = RGBColor(0x15, 0x80, 0x3D)
INK = RGBColor(0x0F, 0x17, 0x2A)
INK2 = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN_HEX = "16A34A"
PANEL_HEX = "F1F5F9"
ACCENT_BG_HEX = "EAF7EF"

# Vendors billed as prepaid credits (booked when purchased).
_PREPAID = {"perplexity", "xai", "grok"}


# ── low-level helpers ───────────────────────────────────────────────────────
def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _cell(cell, text, *, bold=False, color=None, size=9.5, align="left", fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if fill:
        _shade(cell, fill)
    return cell


def _money(x: float) -> str:
    return f"${x:,.2f}"


def _wordmark(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    a = p.add_run("Practice"); a.bold = True; a.font.size = Pt(20); a.font.color.rgb = INK
    b = p.add_run("Rank"); b.bold = True; b.font.size = Pt(20); b.font.color.rgb = GREEN
    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    s = sub.add_run("BY LOST RELIC")
    s.bold = True; s.font.size = Pt(7.5); s.font.color.rgb = MUTED


def _rule(doc, color_hex=GREEN_HEX, size=18):
    """A thin full-width colored rule via a single-cell shaded table."""
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].height = Pt(2)
    _shade(t.rows[0].cells[0], color_hex)
    t.rows[0].cells[0].text = ""
    return t


def _kv_row(table, k, v):
    cells = table.add_row().cells
    _cell(cells[0], k, color=MUTED, bold=True, size=9)
    _cell(cells[1], v, color=INK2, bold=True, size=9, align="right")


# ── document builder ────────────────────────────────────────────────────────
def build(report: dict, out_path: str):
    doc = Document()
    # Base font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.7)
        s.left_margin = s.right_margin = Inches(0.75)

    # Header: wordmark + doc title row
    head = doc.add_table(rows=1, cols=2)
    head.autofit = True
    lc, rc = head.rows[0].cells
    # left: wordmark (build inside the cell)
    lc.text = ""
    p = lc.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    a = p.add_run("Practice"); a.bold = True; a.font.size = Pt(20); a.font.color.rgb = INK
    b = p.add_run("Rank"); b.bold = True; b.font.size = Pt(20); b.font.color.rgb = GREEN
    sp = lc.add_paragraph(); sp.paragraph_format.space_before = Pt(0)
    sr = sp.add_run("BY LOST RELIC"); sr.bold = True; sr.font.size = Pt(7.5); sr.font.color.rgb = MUTED
    # right: eyebrow + title
    rc.text = ""
    ep = rc.paragraphs[0]; ep.alignment = WD_ALIGN_PARAGRAPH.RIGHT; ep.paragraph_format.space_after = Pt(0)
    er = ep.add_run("EXPENSE REPORT · REIMBURSEMENT REQUEST")
    er.bold = True; er.font.size = Pt(8); er.font.color.rgb = GREEN
    tp = rc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; tp.paragraph_format.space_before = Pt(1)
    tr = tp.add_run(report["title"]); tr.bold = True; tr.font.size = Pt(15); tr.font.color.rgb = INK
    pp = rc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; pp.paragraph_format.space_before = Pt(0)
    pr = pp.add_run("Coverage period: " + report["period"]); pr.font.size = Pt(9); pr.font.color.rgb = MUTED
    _no_borders(head)
    _rule(doc)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Meta (2-col key/value)
    meta = doc.add_table(rows=0, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    _kv_row(meta, "Submitted by", report["submitter"])
    _kv_row(meta, "Report date", report["report_date"])
    _kv_row(meta, "Billing email", report["billing_email"])
    _kv_row(meta, "Payment method", report["payment"])
    _kv_row(meta, "Reimburse from", report["reimburse_from"])
    _no_borders(meta)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Summary cards (labels row + amounts row)
    totals = [(_m["label"], sum(i[4] for i in _m["items"])) for _m in report["months"]]
    grand = sum(t for _, t in totals)
    cards = doc.add_table(rows=2, cols=len(totals) + 1)
    for i, (lbl, tot) in enumerate(totals):
        _cell(cards.rows[0].cells[i], lbl.upper(), bold=True, color=MUTED, size=8, fill=PANEL_HEX)
        _cell(cards.rows[1].cells[i], _money(tot), bold=True, color=INK, size=15, fill=PANEL_HEX)
    _cell(cards.rows[0].cells[-1], "TOTAL REIMBURSEMENT", bold=True, color=GREEN_DK, size=8, fill=ACCENT_BG_HEX)
    _cell(cards.rows[1].cells[-1], _money(grand), bold=True, color=GREEN_DK, size=15, fill=ACCENT_BG_HEX)
    _no_borders(cards)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Per-month itemized tables
    headers = ["Category", "Vendor", "Description", "Basis", "Amount"]
    widths = [Inches(1.05), Inches(1.15), Inches(2.85), Inches(0.85), Inches(1.0)]
    for _m in report["months"]:
        sub = sum(i[4] for i in _m["items"])
        h = doc.add_heading(level=2)
        hr = h.add_run(f"{_m['label']}"); hr.font.color.rgb = INK; hr.font.size = Pt(12)
        st = h.add_run(f"      Subtotal {_money(sub)}"); st.font.color.rgb = MUTED; st.font.size = Pt(10); st.bold = True

        t = doc.add_table(rows=1, cols=5)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, hd in enumerate(headers):
            _cell(t.rows[0].cells[i], hd.upper(), bold=True, color=INK2, size=8,
                  align=("right" if i == 4 else "left"), fill=PANEL_HEX)
        for (cat, vendor, desc, basis, amt) in _m["items"]:
            cells = t.add_row().cells
            _cell(cells[0], cat, color=MUTED, size=8.5)
            _cell(cells[1], vendor, bold=True, color=INK, size=9)
            _cell(cells[2], desc, color=INK2, size=9)
            _cell(cells[3], basis, color=MUTED, size=8)
            _cell(cells[4], _money(amt), color=INK, size=9, align="right")
        # subtotal row
        srow = t.add_row().cells
        _cell(srow[0], f"{_m['label']} subtotal", bold=True, color=INK, size=9)
        for c in (srow[1], srow[2], srow[3]):
            _cell(c, "", size=9)
        _cell(srow[4], _money(sub), bold=True, color=INK, size=9.5, align="right", fill=PANEL_HEX)
        _merge(srow[0], srow[3])
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Grand total band
    gt = doc.add_table(rows=1, cols=2)
    _cell(gt.rows[0].cells[0], "TOTAL REIMBURSEMENT REQUESTED", bold=True, color=WHITE, size=11, fill=GREEN_HEX)
    _cell(gt.rows[0].cells[1], _money(grand), bold=True, color=WHITE, size=15, align="right", fill=GREEN_HEX)
    _no_borders(gt)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Notes
    nh = doc.add_paragraph()
    nr = nh.add_run("NOTES & BASIS OF AMOUNTS"); nr.bold = True; nr.font.size = Pt(8); nr.font.color.rgb = INK2
    for note in report["notes"]:
        b = doc.add_paragraph(style="List Bullet")
        r = b.add_run(note); r.font.size = Pt(8.5); r.font.color.rgb = MUTED

    # Sign-off (signature + date filled in)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    so = doc.add_table(rows=1, cols=3)
    _signoff_cell(so.rows[0].cells[0], report.get("signature", ""),
                  "Submitted by — " + report["submitter"].split(",")[0], script=True)
    _signoff_cell(so.rows[0].cells[1], "", "Approved by")
    _signoff_cell(so.rows[0].cells[2], report.get("signed_date", ""), "Date")
    _no_borders(so)

    doc.save(out_path)
    return grand


def _no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        borders.append(e)
    tblPr.append(borders)


def _merge(a, b):
    try:
        a.merge(b)
    except Exception:
        pass


def _set_font_name(run, name):
    """Force a font family across ascii/hAnsi/cs so a script face renders."""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)


def _signoff_cell(cell, value, label, *, script=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if value:
        r = p.add_run(value)
        if script:
            _set_font_name(r, "Brush Script MT")
            r.italic = True
            r.font.size = Pt(18)
            r.font.color.rgb = INK
        else:
            r.font.size = Pt(10.5)
            r.font.color.rgb = INK
    else:
        p.add_run(" ").font.size = Pt(10)
    line = cell.add_paragraph()
    line.paragraph_format.space_before = Pt(0)
    lr = line.add_run("________________________")
    lr.font.size = Pt(9)
    lr.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
    lbl = cell.add_paragraph()
    lbl.paragraph_format.space_before = Pt(0)
    lb = lbl.add_run(label)
    lb.font.size = Pt(8)
    lb.font.color.rgb = MUTED


# ── data: fixed embedded report (May-June 2026) ─────────────────────────────
def default_report() -> dict:
    return {
        "title": "Out-of-Pocket Business Expenses",
        "submitter": "Kody Doherty, Co-founder",
        "report_date": "July 2, 2026",
        "billing_email": "kdoherty@practicerank.ai",
        "payment": "Personal Visa ending 8695",
        "reimburse_from": "PracticeRank (Lost Relic)",
        "period": "May 1 – June 30, 2026",
        "signature": "Kody Doherty",
        "signed_date": "July 2, 2026",
        "months": [
            {"label": "May 2026", "items": [
                ("AI / LLM", "Anthropic", "Claude API — token cost", "Usage", 349.98),
                ("AI / LLM", "Perplexity", "API credit purchase (prepaid)", "Prepaid", 50.00),
                ("AI / LLM", "Google", "Gemini API", "Usage", 21.88),
                ("AI / LLM", "xAI", "Grok API credits (prepaid)", "Prepaid", 25.00),
                ("AI / LLM", "OpenAI", "ChatGPT API — visibility probe", "Usage", 0.05),
                ("Google", "Google Workspace", "Business Standard", "Usage", 21.66),
                ("Infrastructure", "DigitalOcean", "Droplet (shared host)", "Recurring", 24.00),
                ("SEO & Data", "Moz", "Links API (API Starter)", "Recurring", 20.00),
            ]},
            {"label": "June 2026", "items": [
                ("AI / LLM", "Anthropic", "Claude API — token cost", "Usage", 970.81),
                ("AI / LLM", "Anthropic", "Claude web-search tool", "Usage", 104.60),
                ("AI / LLM", "xAI", "Grok API credits — 4× $25 top-ups + $5", "Prepaid", 105.00),
                ("AI / LLM", "OpenAI", "ChatGPT API — mostly web-search fees", "Usage", 20.00),
                ("AI / LLM", "Google", "Gemini API", "Usage", 17.78),
                ("Content COGS", "FATJOE", "Content orders (3)", "COGS", 999.00),
                ("Google", "Google Workspace", "Business Standard", "Usage", 47.32),
                ("Google", "Google Voice", "Telecom", "Usage", 13.12),
                ("Infrastructure", "DigitalOcean", "Droplet (shared host)", "Recurring", 24.00),
                ("SEO & Data", "Moz", "Links API (API Starter)", "Recurring", 20.00),
            ]},
        ],
        "notes": [
            "Recurring — fixed monthly subscription. Usage — metered API/service cost for that "
            "calendar month. Prepaid — API credits, recorded in the month the credits were "
            "purchased. COGS — direct cost of client content delivery.",
            "Anthropic (Claude) ran on the shared organization API key; amounts shown are the "
            "PracticeRank-attributable token/web-search cost (predominantly Opus 4.8 content "
            "generation). Source: Anthropic Console → Analytics → Cost.",
            "Prepaid credits (Perplexity, Grok) are backed by paid invoices/receipts on file "
            "(e.g., Perplexity receipt 2792-0390).",
            "All charges were paid on the submitter's personal Visa ending 8695. Provider console "
            "exports and receipts are retained and available on request.",
        ],
    }


# ── data: pull live from the expense ledger (future months) ─────────────────
def _basis(row) -> str:
    if (row.get("template_key") or "").lower() in _PREPAID:
        return "Prepaid"
    if row.get("recurring"):
        return "Recurring"
    return "Usage"


def report_from_db(db_path: str, months: list[str], period: str) -> dict:
    from geo_agent.db import CustomerDB
    db = CustomerDB(db_path=db_path)
    out_months = []
    for m in months:
        rows = [r for r in db.list_expense_entries(m) if r["amount_usd"] > 0]
        items = []
        for r in sorted(rows, key=lambda r: (r["category"], r["name"])):
            items.append((r["category"], r.get("vendor") or r["name"], r["name"],
                          _basis(r), r["amount_usd"]))
        fj, fn = db.offsite_spend_month(m)
        if fj > 0:
            items.append(("Content COGS", "FATJOE", f"Content orders ({fn})", "COGS", fj))
        y, mo = int(m[:4]), int(m[5:7])
        label = _dt.date(y, mo, 1).strftime("%B %Y")
        out_months.append({"label": label, "items": items})
    db.close()
    rep = default_report()
    rep["months"] = out_months
    rep["period"] = period or ", ".join(months)
    rep["report_date"] = _dt.date.today().strftime("%B %-d, %Y")
    rep["signed_date"] = _dt.date.today().strftime("%B %-d, %Y")
    return rep


def main(argv=None):
    ap = argparse.ArgumentParser(description="Generate PracticeRank expense report .docx")
    ap.add_argument("out", nargs="?", default="PracticeRank-Expense-Report.docx")
    ap.add_argument("--db", help="Path to practicerank.db to pull the ledger live")
    ap.add_argument("--months", help="Comma-separated YYYY-MM list (with --db)")
    ap.add_argument("--period", default="", help="Human coverage period label")
    args = ap.parse_args(argv)

    if args.db and args.months:
        report = report_from_db(args.db, [m.strip() for m in args.months.split(",")], args.period)
    else:
        report = default_report()

    grand = build(report, args.out)
    print(f"Wrote {args.out}  ·  total {_money(grand)}")


if __name__ == "__main__":
    sys.exit(main())
